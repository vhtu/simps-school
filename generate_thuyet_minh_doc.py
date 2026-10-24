"""
Script tạo file Word Thuyết Minh Đề Tài KHKT 2026 hoàn chỉnh
Tên đề tài: SIMPS - Hệ Thống Quản Lý và Đánh Giá Mức Độ Ưu Tiên Sự Cố Cơ Sở Vật Chất Trường Học
Tác giả: Lê Quốc Khánh & Hồ Anh Khoa
GVHD: Phan Thuý Quỳnh
Đơn vị: Trường THPT Ô Môn, Cần Thơ
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        footer = s.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("SIMPS - Dự thi KHKT 2026 | THPT Ô Môn, Cần Thơ | Trang ")
        r_ft.font.name = 'Arial'
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = RGBColor(148, 163, 184)
        
        # Add page number XML
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p_ft._p.append(fldSimple)

def parse_markdown_to_docx(md_path, docx_path):
    print(f"Reading markdown from: {md_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Thiết lập lề chuẩn văn bản học thuật (Top/Bottom 2cm, Left 2.5cm, Right 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

    # Style mặc định
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    style_normal.paragraph_format.line_spacing = 1.25
    style_normal.paragraph_format.space_after = Pt(4)

    add_header_footer(doc)

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []

    def format_inline_runs(paragraph, text, base_font_size=11, base_color=RGBColor(30, 41, 59), italic_base=False):
        # Tách chuỗi theo bold (**text**), italic (*text*), code (`text`)
        pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
        parts = pattern.split(text)
        for part in parts:
            if not part:
                continue
            r = paragraph.add_run()
            r.font.name = 'Arial'
            r.font.size = Pt(base_font_size)
            r.font.color.rgb = base_color
            r.italic = italic_base

            if part.startswith('**') and part.endswith('**'):
                r.text = part[2:-2]
                r.bold = True
            elif part.startswith('*') and part.endswith('*'):
                r.text = part[1:-1]
                r.italic = True
            elif part.startswith('`') and part.endswith('`'):
                r.text = part[1:-1]
                r.font.name = 'Consolas'
                r.font.size = Pt(base_font_size - 1)
                r.font.color.rgb = RGBColor(190, 24, 93)
            else:
                r.text = part

    while i < len(lines):
        line = lines[i]

        # Kiểm tra code block
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                # Thêm bảng cho code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F1F5F9")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run('\n'.join(code_lines))
                r.font.name = 'Consolas'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Kiểm tra bảng Markdown (| Col 1 | Col 2 |)
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            # Xử lý table_lines
            valid_rows = []
            for t_line in table_lines:
                # Bỏ qua dòng separator |---|---|
                cleaned = t_line.strip('|')
                cols = [c.strip() for c in cleaned.split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                    continue
                valid_rows.append(cols)

            if valid_rows:
                num_cols = max(len(r) for r in valid_rows)
                tbl = doc.add_table(rows=len(valid_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Apply borders styling XML
                tblPr = tbl._tbl.tblPr
                borders_xml = parse_xml(
                    f'<w:tblBorders {nsdecls("w")}>\n'
                    f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                    f'  <w:left w:val="none"/>\n'
                    f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>\n'
                    f'  <w:right w:val="none"/>\n'
                    f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
                    f'  <w:insideV w:val="none"/>\n'
                    f'</w:tblBorders>'
                )
                tblPr.append(borders_xml)

                for r_idx, row_data in enumerate(valid_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = tbl.cell(r_idx, c_idx)
                        text = row_data[c_idx] if c_idx < len(row_data) else ""
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        
                        if is_header:
                            set_cell_background(cell, "1E3A8A")
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            format_inline_runs(p, text, base_font_size=10, base_color=RGBColor(255, 255, 255))
                            for run in p.runs:
                                run.bold = True
                        else:
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "FFFFFF")
                            else:
                                set_cell_background(cell, "F8FAFC")
                            p = cell.paragraphs[0]
                            format_inline_runs(p, text, base_font_size=9.5, base_color=RGBColor(30, 41, 59))
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(6)
            continue

        # Tiêu đề Markdown
        stripped = line.strip()
        if stripped.startswith('# '):
            # Title chính
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            format_inline_runs(p, stripped[2:], base_font_size=18, base_color=RGBColor(30, 58, 138))
            for run in p.runs:
                run.bold = True
        elif stripped.startswith('## '):
            # Phần lớn (H1)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            # Thêm đường viền trang trí hoặc màu nổi
            format_inline_runs(p, stripped[3:], base_font_size=13.5, base_color=RGBColor(30, 58, 138))
            for run in p.runs:
                run.bold = True
        elif stripped.startswith('### '):
            # Mục con (H2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            format_inline_runs(p, stripped[4:], base_font_size=12, base_color=RGBColor(14, 116, 144))
            for run in p.runs:
                run.bold = True
        elif stripped.startswith('#### '):
            # Mục nhỏ hơn (H3)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, stripped[5:], base_font_size=11, base_color=RGBColor(71, 85, 105))
            for run in p.runs:
                run.bold = True
        elif stripped.startswith('> '):
            # Trích dẫn / Callout
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "EFF6FF")
            # Left border blue
            borders_xml = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>\n'
                f'  <w:top w:val="none"/>\n'
                f'  <w:bottom w:val="none"/>\n'
                f'  <w:right w:val="none"/>\n'
                f'</w:tcBorders>'
            )
            cell._tc.get_or_add_tcPr().append(borders_xml)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            p = cell.paragraphs[0]
            format_inline_runs(p, stripped[2:], base_font_size=10.5, base_color=RGBColor(30, 64, 175), italic_base=True)
            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            format_inline_runs(p, stripped[2:], base_font_size=11, base_color=RGBColor(30, 41, 59))
        elif re.match(r'^\d+\.\s+', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text_after_num = re.sub(r'^\d+\.\s+', '', stripped)
            format_inline_runs(p, text_after_num, base_font_size=11, base_color=RGBColor(30, 41, 59))
        elif stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Dòng phân cách ngang
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'CBD5E1')
            pbdr.append(bottom)
            pPr.append(pbdr)
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            format_inline_runs(p, stripped, base_font_size=11, base_color=RGBColor(30, 41, 59))

        i += 1

    print(f"Saving to Word docx at: {docx_path}")
    doc.save(docx_path)
    print("Export completed successfully!")

if __name__ == '__main__':
    md_file = r"D:\KHKT\THUYET_MINH_HE_THONG_SIMPS_KHKT2026.md"
    docx_file = r"D:\KHKT\Thuyet_minh_he_thong_SIMPS_KHKT2026.docx"
    parse_markdown_to_docx(md_file, docx_file)
