"""
Script tạo file Word (.docx) Hướng dẫn Sử dụng Hệ thống SIMPS
Dành cho đề tài KHKT 2026 - Trường THPT Ô Môn, Cần Thơ.
Học sinh thực hiện: Lê Quốc Khánh và Hồ Anh Khoa.
"""

import os
from datetime import datetime
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def create_manual():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(4)

    # Header
    p_top = doc.add_paragraph()
    r1 = p_top.add_run("SỞ GD&ĐT THÀNH PHỐ CẦN THƠ • TRƯỜNG THPT Ô MÔN\nCUỘC THI KHOA HỌC KỸ THUẬT DÀNH CHO HỌC SINH TRUNG HỌC NĂM 2026\n")
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(100, 116, 139)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("TÀI LIỆU HƯỚNG DẪN SỬ DỤNG HỆ THỐNG SIMPS\nĐỀ TÀI: NGHIÊN CỨU, PHÁT TRIỂN HỆ THỐNG TIẾP NHẬN VÀ TỰ ĐỘNG XẾP HẠNG ƯU TIÊN XỬ LÝ SỰ CỐ CƠ SỞ VẬT CHẤT TRƯỜNG HỌC DỰA TRÊN THUẬT TOÁN ĐA TIÊU CHÍ (SIMPS)")
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Học sinh thực hiện: Lê Quốc Khánh và Hồ Anh Khoa\n"
                          "Giáo viên hướng dẫn (GVHD): Phan Thuý Quỳnh\n"
                          "Đơn vị: Trường THPT Ô Môn, TP Cần Thơ\n"
                          "Hệ thống trực tuyến: https://simps-school.vercel.app")
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.paragraph_format.space_after = Pt(18)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(14, 116, 144)
        return p

    # Phần 1
    add_h1("PHẦN 1: GIỚI THIỆU TỔNG QUAN")
    doc.add_paragraph(
        "Hệ thống SIMPS là phần mềm hỗ trợ quản trị và giải quyết nhanh các hư hỏng cơ sở vật chất (CSVC) trong trường học, "
        "ứng dụng thuật toán chấm điểm ưu tiên đa tiêu chí tuyến tính kết hợp cơ chế ghi đè an toàn (Safety Override):\n"
        "P = w1*S + w2*I + w3*R + w4*T + w5*E (với S: Nghiêm trọng, I: Số người ảnh hưởng, R: Rủi ro vị trí, T: Thời gian chờ, E: Tác động học tập).\n"
        "Bộ trọng số được kiểm chứng khoa học theo phương pháp AHP của Saaty với CR = 0.042 < 0.10."
    )

    # Phần 2
    add_h1("PHẦN 2: HƯỚNG DẪN DÀNH CHO HỌC SINH (HỌC ĐƯỜNG)")
    add_h2("2.1. Đăng nhập hệ thống")
    doc.add_paragraph("• Học sinh truy cập địa chỉ: https://simps-school.vercel.app, nhập tài khoản (ví dụ: hs01 / 123456) và bấm 'Đăng Nhập Ngay'.")

    add_h2("2.2. Báo sự cố 0-Touch qua Tem Mã QR Phòng Học")
    doc.add_paragraph(
        "• Khi phát hiện hỏng hóc trong lớp: Mở Zalo hoặc Camera điện thoại quét Tem Mã QR dán tại cửa phòng học.\n"
        "• Hệ thống tự động nhận diện chính xác phòng học và tầng lầu, học sinh chỉ cần chọn loại hư hỏng và bấm 'Gửi Báo Cáo' trong vòng 20 giây."
    )

    add_h2("2.3. Nghiệm thu & Xác nhận hoàn thành")
    doc.add_paragraph(
        "• Sau khi kỹ thuật viên sửa xong thiết bị, sự cố chuyển sang trạng thái 'Đã khắc phục'.\n"
        "• Học sinh vào tab 'Báo Cáo Của Tôi', kiểm tra thiết bị thực tế và bấm nút '✅ Xác nhận' để đóng sự cố."
    )

    # Phần 3
    add_h1("PHẦN 3: HƯỚNG DẪN DÀNH CHO TỔ KỸ THUẬT & CSVC (STAFF)")
    add_h2("3.1. Hàng đợi ưu tiên tự động (Core Queue)")
    doc.add_paragraph("• Danh sách sự cố được tự động sắp xếp theo điểm P giảm dần. Kỹ thuật viên giải quyết sự cố từ trên xuống dưới, đảm bảo tính công bằng và loại bỏ hoàn toàn cảm tính.")

    add_h2("3.2. Cảnh báo an toàn đặc biệt (Safety Override)")
    doc.add_paragraph("• Khi phát hiện nguy cơ cao (chập điện, nổ bình gas, rơi vỡ...) với P = 999, màn hình hiển thị Banner đỏ nhấp nháy, yêu cầu can thiệp trong vòng dưới 15 phút.")

    add_h2("3.3. Cập nhật tiến độ sửa chữa")
    doc.add_paragraph("• Bấm nút '⚙️ Thao tác' để nhận việc, chuyển trạng thái từ 'Đang xử lý' sang 'Đã khắc phục' và ghi lại giải pháp linh kiện thay thế.")

    # Phần 4
    add_h1("PHẦN 4: HƯỚNG DẪN DÀNH CHO QUẢN TRỊ VIÊN & BAN GIÁM HIỆU (ADMIN)")
    add_h2("4.1. Bảng điều khiển phân tích & Khuyến nghị bảo trì dự đoán")
    doc.add_paragraph(
        "• Giám sát thời gian thực qua 5 biểu đồ Canvas Retina: Phân bổ hư hỏng, Điểm nóng phòng học, Tần suất theo tháng, Thời gian xử lý trung bình và Hiệu suất kỹ thuật viên.\n"
        "• Hệ thống tự động phân tích tần suất hỏng lặp lại để đưa ra khuyến nghị 'ĐẠI TU TOÀN DIỆN' hoặc 'THAY THẾ THIẾT BỊ MỚI'."
    )

    add_h2("4.2. In hàng loạt tem QR phòng học & Cấu hình AHP")
    doc.add_paragraph(
        "• Vào mục 'Địa Điểm & Mức Rủi Ro': Bấm nút '🖨️ In Hàng Loạt Tem QR Phòng Học' để tạo bản in A4 chứa toàn bộ tem dán cửa lớp cho cả trường.\n"
        "• Vào mục 'Trọng Số & Bộ Giả Lập P': Điều chỉnh trọng số w1-w5 và kiểm chứng ma trận Saaty."
    )

    add_h2("4.3. Quản lý Địa điểm, Loại sự cố và Tài khoản")
    doc.add_paragraph(
        "• Hỗ trợ đầy đủ chức năng Thêm mới, Sửa thông tin, Đổi mật khẩu người dùng và Xóa an toàn cho cả Địa điểm, Loại sự cố và Tài khoản."
    )

    add_h2("4.4. Cấu hình Bắn thông báo tức thì (Zalo Webhook & Telegram Bot)")
    doc.add_paragraph(
        "• Vào mục '📲 Bắn Thông Báo Zalo & Telegram': Dán Telegram Bot Token và Chat ID (hoặc Zalo Webhook URL), bấm 'Gửi thử tin nhắn' để kiểm tra kết nối.\n"
        "• Khi có sự cố khẩn cấp, hệ thống tự động bắn tin nhắn báo động ra điện thoại của Ban Giám hiệu và Tổ Kỹ thuật."
    )

    # Bảng tài khoản mẫu
    add_h1("PHẦN 5: DANH SÁCH TÀI KHOẢN MẪU DÙNG THỬ (DEMO)")
    t = doc.add_table(rows=5, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Vai trò", "Tên đăng nhập", "Mật khẩu", "Mục đích trải nghiệm"]
    for idx, h in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.text = h
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

    data = [
        ("Quản trị viên (Admin)", "admin", "admin123", "Bảng phân tích, cấu hình AHP, in QR, webhook Zalo/Telegram"),
        ("Kỹ thuật Điện (Staff)", "nhanvien_dien", "123456", "Hàng đợi ưu tiên P, cảnh báo Safety Override, cập nhật sửa chữa"),
        ("Kỹ thuật CSVC (Staff)", "nhanvien_nuoc", "123456", "Xử lý sự cố nước, bàn ghế, cửa kính, thiết bị trường học"),
        ("Học sinh (Student)", "hs01 đến hs20", "123456", "Quét mã QR báo sự cố mới, theo dõi tiến trình, nghiệm thu xác nhận")
    ]
    for r_idx, row_data in enumerate(data):
        row = t.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row[c_idx].text = val
            p = row[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs: r.font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
            set_cell_background(row[c_idx], bg)

    # Footer note
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_time = p_end.add_run(f"Tài liệu xuất bản tự động từ Hệ thống SIMPS • Ngày {datetime.now().strftime('%d/%m/%Y')}")
    r_time.font.italic = True
    r_time.font.size = Pt(9.5)
    r_time.font.color.rgb = RGBColor(148, 163, 184)

    output_path = r"D:\KHKT\Huong_dan_su_dung_he_thong_SIMPS.docx"
    doc.save(output_path)
    print(f"Generated docx successfully: {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    create_manual()
